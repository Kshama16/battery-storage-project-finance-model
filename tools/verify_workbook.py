"""Checks a recalculated copy of the workbook against the Python reference model.

The workbook must have been recalculated by a spreadsheet engine (Excel or
LibreOffice) so that cached values are stored. For each scenario, set
Scenarios!C5, recalculate, and compare; or pass one file per scenario.

    python tools/verify_workbook.py recalc_base.xlsx recalc_downside.xlsx recalc_upside.xlsx
"""
import sys

from openpyxl import load_workbook

from reference_model import Case, required_capacity_price, run

METRICS = [
    # defined name, reference key, tolerance (absolute)
    ("TotalCapex", "total_capex", 0.5),
    ("SrAmt", "senior", 0.5),
    ("MzAmt", "mezz", 0.5),
    ("Gearing", "gearing", 1e-5),
    ("WACC", "wacc", 1e-6),
    ("ProjIRR", "project_irr", 1e-6),
    ("EqIRR", "equity_irr", 1e-6),
    ("ProjNPV", "project_npv", 0.5),
    ("EqNPV", "equity_npv", 0.5),
    ("MinDSCRSr", "min_dscr_senior", 1e-4),
    ("MinDSCRComb", "min_dscr_combined", 1e-4),
    ("AvgDSCRSr", "avg_dscr_senior", 1e-4),
    ("Payback", "payback_year", 0),
    ("ConvYear", "conversion_year", 0),
]


def named_value(wb, name):
    dn = wb.defined_names[name]
    (sheet, ref), = dn.destinations
    return wb[sheet][ref.replace("$", "")].value


def check(path):
    wb = load_workbook(path, data_only=True)
    sel = named_value(wb, "ScenarioSel")
    scen_name = ["Base", "Downside", "Upside"][int(sel) - 1]
    ref = run(Case.scenario(scen_name))
    print(f"\n{path}  ({scen_name})")
    ok = True
    for name, key, tol in METRICS:
        x, y = named_value(wb, name), ref[key]
        good = x is not None and not isinstance(x, str) and abs(x - y) <= tol
        ok &= good
        print(f"  {'OK ' if good else 'BAD'} {name:12s} workbook {x!s:>22}  reference {y!s:>22}")

    all_checks = wb["Checks"]["C4"].value
    print(f"  {'OK ' if all_checks is True else 'BAD'} Checks!C4 = {all_checks}")
    ok &= all_checks is True

    if scen_name == "Base":
        ok &= check_data_tables(wb["Scenarios"])
        req = named_value(wb, "RequiredPrice")
        exact = required_capacity_price()
        # The workbook interpolates linearly on a 2%-step data table, so allow a few cents
        good = isinstance(req, (int, float)) and abs(req - exact) < 0.05
        ok &= good
        print(f"  {'OK ' if good else 'BAD'} RequiredPrice workbook {req}  reference {exact:.4f}")
    return ok


# One-variable tables on the Scenarios sheet: title prefix -> how a flex value maps onto the base case
FLEX = {
    "Scenario comparison": lambda v: Case.scenario(["Base", "Downside", "Upside"][int(v) - 1]),
    "Capex sensitivity": lambda v: Case(capex_adj=v),
    "Capacity price sensitivity": lambda v: Case(price_adj=v),
    "Construction delay": lambda v: Case(delay=int(v)),
    "Interest rates": lambda v: Case(rate_bps=v),
}
TABLE_OUTPUTS = {"Equity IRR": "equity_irr", "Project IRR": "project_irr", "Project NPV": "project_npv",
                 "Equity NPV": "equity_npv", "Min senior DSCR": "min_dscr_senior"}


def close(x, y):
    return isinstance(x, (int, float)) and abs(x - y) <= max(1e-6, 1e-6 * abs(y))


def check_data_tables(ws):
    """Compares every computed data-table cell with the reference model. Blank cells are
    reported but not failed: LibreOffice only evaluates the first two columns of a
    multi-column one-variable table (Excel evaluates all of them)."""
    ok, n_checked, n_blank = True, 0, 0
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 2).value
        if not isinstance(label, str):
            continue
        key = next((k for k in FLEX if label.startswith(k)), None)
        if key:
            hdr = row + 1
            outputs = {c: TABLE_OUTPUTS.get(ws.cell(hdr, c).value) for c in range(3, 9)}
            r = hdr + 2
            while isinstance(ws.cell(r, 2).value, (int, float)):
                ref = run(FLEX[key](ws.cell(r, 2).value))
                for c, out in outputs.items():
                    x = ws.cell(r, c).value
                    if out is None:
                        continue
                    if x is None:
                        n_blank += 1
                        continue
                    n_checked += 1
                    if not close(x, ref[out]):
                        ok = False
                        print(f"  BAD {label[:30]} {ws.cell(r, 2).value} {out}: workbook {x} reference {ref[out]}")
                r += 1
        elif label.startswith("Equity IRR: capex flex"):
            corner = row + 1
            for i in range(1, 6):
                for j in range(1, 6):
                    capex, price = ws.cell(corner + i, 2).value, ws.cell(corner, 2 + j).value
                    x, y = ws.cell(corner + i, 2 + j).value, run(Case(capex_adj=capex, price_adj=price))["equity_irr"]
                    n_checked += 1
                    if not close(x, y):
                        ok = False
                        print(f"  BAD 2-D table capex {capex} price {price}: workbook {x} reference {y}")
    print(f"  {'OK ' if ok else 'BAD'} Data tables: {n_checked} cells match reference"
          + (f" ({n_blank} blank cells not evaluated by this engine)" if n_blank else ""))
    return ok


if __name__ == "__main__":
    results = [check(p) for p in sys.argv[1:]]
    print("\nALL MATCH" if all(results) else "\nMISMATCHES FOUND")
    sys.exit(0 if all(results) else 1)
