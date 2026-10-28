"""Builds model/Battery_Storage_Project_Finance_Model.xlsx with live Excel formulas.

Every calculation in the workbook is an Excel formula driven by the blue input
cells; this script only lays the model out. Inputs come from assumptions.py.

    python tools/build_workbook.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import DataTableFormula

from assumptions import ASSETS, FINANCING, GLOBAL, MACRS, MODEL_START, MODEL_YEARS, SCENARIOS

OUT = Path(__file__).resolve().parent.parent / "model" / "Battery_Storage_Project_Finance_Model.xlsx"

N_MONTHS = MODEL_YEARS * 12
FIRST = 5  # first timeline column (E) on Monthly and Annual

# ---------- styles ----------
FONT = "Calibri"
NAVY, TEAL, LIGHT, INPUT_FILL, GREY = "1F3B4D", "2C7A7B", "E8F1F2", "FFF7D6", "F2F2F2"
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=10, italic=True, color="666666")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_sect = Font(name=FONT, size=11, bold=True, color=NAVY)
f_bold = Font(name=FONT, size=10, bold=True)
f_norm = Font(name=FONT, size=10)
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="00753A")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_sect = PatternFill("solid", fgColor=LIGHT)
fill_input = PatternFill("solid", fgColor=INPUT_FILL)
fill_total = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="BFBFBF")
top_border = Border(top=Side(style="thin", color="000000"))
box = Border(left=thin, right=thin, top=thin, bottom=thin)

K = '#,##0;(#,##0);"–"'           # $k
K1 = '#,##0.0;(#,##0.0);"–"'
PCT = '0.0%;(0.0%);"–"'
PCT2 = '0.00%;(0.00%);"–"'
X2 = '0.00"x"'
DATE = "mmm-yy"
USD2 = '$#,##0.00'

wb = Workbook()
names = {}


def define(name, sheet, ref):
    names[name] = f"'{sheet.title}'!{ref}"
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet.title}'!{_abs(ref)}")


def _abs(ref):
    parts = []
    for p in ref.split(":"):
        letters = "".join(c for c in p if c.isalpha())
        digits = "".join(c for c in p if c.isdigit())
        parts.append(f"${letters}${digits}")
    return ":".join(parts)


def put(ws, ref, value, font=f_norm, fmt=None, fill=None, bold=False, align=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=font.size, bold=bold or font.bold, italic=font.italic, color=font.color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = border
    return c


def inp(ws, ref, value, fmt=None):
    return put(ws, ref, value, font=f_input, fmt=fmt, fill=fill_input, border=box)


def header_row(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill_head
        cell.font = f_head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def section(ws, row, text, last_col=8):
    for c in range(2, last_col + 1):
        ws.cell(row=row, column=c).fill = fill_sect
    put(ws, f"B{row}", text, font=f_sect)


def title(ws, text, sub):
    put(ws, "B2", text, font=f_title)
    put(ws, "B3", sub, font=f_sub)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2


# =====================================================================
# Sheets (created up-front so tab order is fixed)
# =====================================================================
cover = wb.active
cover.title = "Cover"
summary = wb.create_sheet("Summary")
inputs = wb.create_sheet("Inputs")
scen = wb.create_sheet("Scenarios")
monthly = wb.create_sheet("Monthly")
annual = wb.create_sheet("Annual")
log = wb.create_sheet("Assumption Log")
checks = wb.create_sheet("Checks")

for ws, color in [(cover, NAVY), (summary, NAVY), (inputs, "C9A227"), (scen, "C9A227"),
                  (monthly, TEAL), (annual, TEAL), (log, "7F7F7F"), (checks, "7F7F7F")]:
    ws.sheet_properties.tabColor = color

# =====================================================================
# Inputs
# =====================================================================
title(inputs, "Inputs", "Blue cells on yellow are inputs. Black cells are formulas. $ figures in thousands ($k) unless stated.")
inputs.column_dimensions["B"].width = 38
inputs.column_dimensions["C"].width = 12
inputs.column_dimensions["D"].width = 14
inputs.column_dimensions["E"].width = 62

GLOBAL_ROWS = [
    # name, label, value, fmt, unit, note
    ("ModelStart", "Model start / valuation date", MODEL_START, "dd-mmm-yyyy", "date", "Monthly timeline starts here; NPVs are discounted to this date"),
    ("CapPrice", "Capacity price (2027 $)", GLOBAL["capacity_price"], USD2, "$/kW-month", "Paid on accredited MW. Use Goal Seek or Scenarios!required price to back-solve"),
    ("PriceEsc", "Capacity price & margin escalation", GLOBAL["price_escalation"], PCT, "% / yr", ""),
    ("OMEsc", "O&M, lease & G&A escalation", GLOBAL["om_escalation"], PCT, "% / yr", ""),
    ("Contingency", "Owner's costs & contingency", GLOBAL["contingency"], PCT, "% of EPC", "Applied on top of EPC unit costs"),
    ("Insurance", "Insurance", GLOBAL["insurance"], PCT2, "% capex / yr", ""),
    ("PropTax", "Property tax", GLOBAL["property_tax"], PCT2, "% capex / yr", ""),
    ("GACost", "Portfolio G&A and asset management", GLOBAL["ga_cost"], K, "$k / yr", "Starts at first COD, escalates with O&M"),
    ("TaxRate", "Blended income tax rate", GLOBAL["tax_rate"], PCT, "%", "Losses assumed usable by sponsor (negative tax = benefit)"),
    ("ITCRate", "Investment tax credit (storage only)", GLOBAL["itc_rate"], PCT, "% eligible capex", "Monetized via transfer; cash received the year after COD"),
    ("ITCTransfer", "ITC transfer price", GLOBAL["itc_transfer_price"], USD2, "$ per $1 credit", ""),
    ("ITCBasisRed", "Depreciable basis reduction", GLOBAL["itc_basis_reduction"], PCT, "% of ITC", ""),
]
FIN_ROWS = [
    ("SrRateBase", "Senior debt interest rate", FINANCING["senior_rate"], PCT2, "% / yr", "Live rate adds Scenarios rate adjustment"),
    ("SrDSCR", "Senior target DSCR", FINANCING["senior_dscr"], X2, "x", "Debt service sculpted to CFADS / target DSCR"),
    ("SrMaxGear", "Senior max gearing", FINANCING["senior_max_gearing"], PCT, "% of capex", "Cap on senior sizing"),
    ("SrFee", "Senior upfront fee", FINANCING["senior_fee"], PCT2, "% of amount", ""),
    ("MzRateBase", "Mezzanine interest rate", FINANCING["mezz_rate"], PCT2, "% / yr", ""),
    ("CombDSCR", "Combined (senior + mezz) target DSCR", FINANCING["combined_dscr"], X2, "x", "Mezz takes the gap between senior and combined DSCR"),
    ("TotMaxGear", "Total max gearing", FINANCING["total_max_gearing"], PCT, "% of capex", ""),
    ("MzFee", "Mezzanine upfront fee", FINANCING["mezz_fee"], PCT2, "% of amount", ""),
    ("Tenor", "Debt tenor (both tranches)", FINANCING["tenor"], "0", "years", "From term conversion (latest COD)"),
    ("CostEquity", "Cost of equity / target equity IRR", FINANCING["cost_of_equity"], PCT, "%", "Hurdle for required capacity price"),
]

r = 5
for heading, rows in [("General & market", GLOBAL_ROWS), ("Financing", FIN_ROWS)]:
    section(inputs, r, heading, 5)
    r += 1
    for h, txt in zip("BCDE", ["Assumption", "Value", "Unit", "Note"]):
        put(inputs, f"{h}{r}", txt)
    header_row(inputs, r, 2, 5)
    r += 1
    for name, label, value, fmt, unit, note in rows:
        put(inputs, f"B{r}", label)
        inp(inputs, f"C{r}", value, fmt)
        put(inputs, f"D{r}", unit, font=f_sub)
        put(inputs, f"E{r}", note, font=f_sub)
        define(name, inputs, f"C{r}")
        r += 1
    r += 1

# Asset table
section(inputs, r, "Asset portfolio", 26)
r += 1
ASSET_COLS = [
    # header, key, fmt, width
    ("Asset", "name", None, 10), ("Technology", "tech", None, 11), ("MW", "mw", "#,##0", 7),
    ("Duration (h)", "hours", "0", 8), ("EPC unit cost", "cost", "#,##0", 9), ("Cost unit", "unit", None, 8),
    ("Build (months)", "build", "0", 8), ("COD (base)", "cod", DATE, 9), ("Capacity accreditation", "accred", "0%", 11),
    ("Degradation / yr", "degr", PCT, 10), ("Fixed O&M $/kW-yr", "fom", "#,##0", 9),
    ("Energy & ancillary margin $/kW-yr", "margin", "#,##0", 11), ("Site lease $k/yr", "lease", "#,##0", 8),
    ("ITC eligible (1/0)", "itc", "0", 8), ("MACRS life (5/15)", "macrs", "0", 8), ("Useful life (yrs)", "life", "0", 8),
]
CALC_COLS = [
    ("$/kW (EPC)", '=F{r}*IF(G{r}="$/kWh",E{r},1)', "#,##0", 9),
    ("MWh", "=D{r}*E{r}", "#,##0", 8),
    ("Total capex $k", "=D{r}*R{r}*(1+Contingency)*(1+CapexAdj)", K, 11),
    ("COD (live, incl. delay)", "=EDATE(I{r},Delay)", DATE, 10),
    ("COD year", "=YEAR(U{r})", "0", 7),
    ("Capex start", "=EDATE(I{r},-H{r})", DATE, 9),
    ("ITC $k", "=T{r}*ITCRate*O{r}", K, 9),
    ("Tax basis $k", "=T{r}*(1-ITCBasisRed*ITCRate*O{r})", K, 10),
    ("Life end", "=EDATE(U{r},12*Q{r})", DATE, 9),
]
hdr = r
for i, (h, _, _, w) in enumerate(ASSET_COLS + [(c[0], None, None, c[3]) for c in CALC_COLS]):
    c = 2 + i
    put(inputs, f"{col(c)}{hdr}", h)
    if c > 5:
        inputs.column_dimensions[col(c)].width = max(inputs.column_dimensions[col(c)].width or 0, w)
header_row(inputs, hdr, 2, 2 + len(ASSET_COLS) + len(CALC_COLS) - 1)
inputs.row_dimensions[hdr].height = 42
put(inputs, f"{col(2 + len(ASSET_COLS))}{hdr - 1}", "Calculated (do not edit) →", font=f_sub)

ASSET_ROW = {}
r = hdr + 1
for a in ASSETS:
    ASSET_ROW[a["name"]] = r
    for i, (_, key, fmt, _) in enumerate(ASSET_COLS):
        inp(inputs, f"{col(2 + i)}{r}", a[key], fmt)
    for i, (_, formula, fmt, _) in enumerate(CALC_COLS):
        put(inputs, f"{col(2 + len(ASSET_COLS) + i)}{r}", formula.format(r=r), fmt=fmt)
    r += 1
A1, A2 = hdr + 1, r - 1
put(inputs, f"B{r}", "Portfolio", bold=True)
for c_ in ["D", "S", "T", "X", "Y"]:
    put(inputs, f"{c_}{r}", f"=SUM({c_}{A1}:{c_}{A2})", fmt=K, bold=True, border=top_border)
define("TotalCapexInputs", inputs, f"T{r}")
put(inputs, f"B{r + 1}", "EPC unit cost is $/kWh for storage (× duration) and $/kW for generation.", font=f_sub)

dv_tech = DataValidation(type="list", formula1='"$/kWh,$/kW"', allow_blank=False)
inputs.add_data_validation(dv_tech)
dv_tech.add(f"G{A1}:G{A2}")
dv_macrs = DataValidation(type="list", formula1='"5,15"', allow_blank=False)
inputs.add_data_validation(dv_macrs)
dv_macrs.add(f"P{A1}:P{A2}")
dv_itc = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
inputs.add_data_validation(dv_itc)
dv_itc.add(f"O{A1}:O{A2}")

# MACRS table
r += 3
section(inputs, r, "MACRS depreciation (half-year convention)", 5)
r += 1
for h, txt in zip("BCD", ["Tax year", "5-year", "15-year"]):
    put(inputs, f"{h}{r}", txt)
header_row(inputs, r, 2, 4)
m1 = r + 1
for k in range(16):
    put(inputs, f"B{m1 + k}", k + 1)
    inp(inputs, f"C{m1 + k}", MACRS[5][k] if k < 6 else 0, "0.00%")
    inp(inputs, f"D{m1 + k}", MACRS[15][k], "0.00%")
define("MACRS_5", inputs, f"C{m1}:C{m1 + 15}")
define("MACRS_15", inputs, f"D{m1}:D{m1 + 15}")
put(inputs, f"B{m1 + 16}", "Total", bold=True)
put(inputs, f"C{m1 + 16}", f"=SUM(C{m1}:C{m1 + 15})", fmt="0.00%", bold=True, border=top_border)
put(inputs, f"D{m1 + 16}", f"=SUM(D{m1}:D{m1 + 15})", fmt="0.00%", bold=True, border=top_border)
inputs.freeze_panes = "C5"

# =====================================================================
# Scenarios: live drivers + sensitivity data tables (Excel requires the
# data-table input cells to sit on the same sheet as the tables)
# =====================================================================
title(scen, "Scenarios & Sensitivities",
      "Pick the live case in C5. Sensitivity flex (column G) layers on top of the scenario. Tables below are Excel Data Tables (What-If) and refresh on recalc (F9).")
for c_, w in zip("BCDEFGHIJKLMNO", [36, 12, 12, 12, 13, 13, 12, 3, 12, 12, 13, 3, 22, 12]):
    scen.column_dimensions[c_].width = w

put(scen, "B5", "Active scenario (1 Base, 2 Downside, 3 Upside)", bold=True)
inp(scen, "C5", 1, "0")
define("ScenarioSel", scen, "C5")
put(scen, "D5", '=INDEX(C8:E8,C5)', bold=True)
dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
scen.add_data_validation(dv)
dv.add("C5")

for c_, txt in zip("BCDEFGH", ["Driver", "", "", "", "Active scenario", "Sensitivity flex", "Live value"]):
    put(scen, f"{c_}7", txt)
header_row(scen, 7, 2, 8)
put(scen, "B8", "Scenario name", font=f_sub)
for c_, n in zip("CDE", SCENARIOS):
    inp(scen, f"{c_}8", n)
DRIVERS = [
    ("CapexAdj", "Capex adjustment", "capex", PCT),
    ("PriceAdj", "Capacity price adjustment", "price", PCT),
    ("Delay", "COD delay (months)", "delay", "0"),
    ("RateAdj", "Interest rate adjustment (bps)", "rate", "0"),
    ("OMAdj", "Fixed O&M adjustment", "om", PCT),
]
DRIVER_ROW = {}
for i, (name, label, key, fmt) in enumerate(DRIVERS):
    rr = 9 + i
    DRIVER_ROW[name] = rr
    put(scen, f"B{rr}", label)
    for c_, n in zip("CDE", SCENARIOS):
        inp(scen, f"{c_}{rr}", SCENARIOS[n][key], fmt)
    put(scen, f"F{rr}", f"=CHOOSE(ScenarioSel,C{rr},D{rr},E{rr})", fmt=fmt)
    inp(scen, f"G{rr}", 0, fmt)
    put(scen, f"H{rr}", f"=F{rr}+G{rr}", fmt=fmt, bold=True)
    define(name, scen, f"H{rr}")
put(scen, "B14", "Delay pushes every COD (and term conversion) back; construction spend timing is unchanged.", font=f_sub)

# =====================================================================
# Monthly
# =====================================================================
title(monthly, "Monthly operating cash flows", "Asset-level operations and capex by month, $k. Timeline columns are formulas off Inputs!ModelStart.")
monthly.column_dimensions["B"].width = 34
monthly.column_dimensions["C"].width = 12
monthly.column_dimensions["D"].width = 12
LAST_M = FIRST + N_MONTHS - 1
LM = col(LAST_M)
for c in range(FIRST, LAST_M + 1):
    monthly.column_dimensions[col(c)].width = 9

put(monthly, "C5", "Unit", bold=True)
put(monthly, "D5", "Total", bold=True)
ROW = {}


def mrow(r_, label, unit, formula, fmt=K, total=True, bold=False, key=None, font=f_norm):
    put(monthly, f"B{r_}", label, bold=bold, font=font)
    put(monthly, f"C{r_}", unit, font=f_sub)
    for c in range(FIRST, LAST_M + 1):
        L, P = col(c), col(c - 1)
        put(monthly, f"{L}{r_}", formula(L, P), fmt=fmt, bold=bold, font=font)
    if total:
        put(monthly, f"D{r_}", f"=SUM({col(FIRST)}{r_}:{LM}{r_})", fmt=fmt, bold=True)
    if key:
        ROW[key] = r_


mrow(5, "Month starting", "", lambda L, P: "=ModelStart" if L == col(FIRST) else f"=EDATE({P}5,1)", fmt=DATE, total=False, bold=True)
header_row(monthly, 5, FIRST, LAST_M)
monthly.cell(row=5, column=2).font = f_bold
mrow(6, "Year", "", lambda L, P: f"=YEAR({L}5)", fmt="0", total=False)
mrow(7, "Years from model start", "", lambda L, P: f"={L}6-YEAR(ModelStart)", fmt="0", total=False)
mrow(8, "Capacity price (live)", "$/kW-mo", lambda L, P: f"=CapPrice*(1+PriceAdj)*(1+PriceEsc)^{L}7", fmt=USD2, total=False)
mrow(9, "Price escalation index", "x", lambda L, P: f"=(1+PriceEsc)^{L}7", fmt="0.000", total=False)
mrow(10, "O&M escalation index", "x", lambda L, P: f"=(1+OMEsc)^{L}7", fmt="0.000", total=False)

r = 12
ASSET_BLOCK = {}
for a in ASSETS:
    ir = ASSET_ROW[a["name"]]
    I = lambda c_: f"Inputs!${c_}${ir}"
    section(monthly, r, f"{a['name']} — {a['tech']} {a['mw']} MW" + (f" / {a['mw'] * a['hours']} MWh" if a["hours"] else ""), 4)
    b = r + 1
    ASSET_BLOCK[a["name"]] = b
    mrow(b, "Operating flag", "1/0", lambda L, P: f"=IF(AND({L}$5>={I('U')},{L}$5<{I('Z')}),1,0)", fmt="0", total=False)
    mrow(b + 1, "Operating year", "#", lambda L, P: f"=IF({L}{b}=1,INT(((YEAR({L}$5)-YEAR({I('U')}))*12+MONTH({L}$5)-MONTH({I('U')}))/12),0)", fmt="0", total=False)
    mrow(b + 2, "Accredited capacity (after degradation)", "MW", lambda L, P: f"={L}{b}*{I('D')}*{I('J')}*(1-{I('K')})^{L}{b + 1}", fmt=K1, total=False)
    mrow(b + 3, "Capacity revenue", "$k", lambda L, P: f"={L}{b + 2}*{L}$8")
    mrow(b + 4, "Energy & ancillary margin", "$k", lambda L, P: f"={L}{b + 2}*{I('M')}/12*{L}$9")
    mrow(b + 5, "Fixed O&M", "$k", lambda L, P: f"={L}{b}*{I('D')}*{I('L')}/12*{L}$10*(1+OMAdj)")
    mrow(b + 6, "Insurance & property tax", "$k", lambda L, P: f"={L}{b}*{I('T')}*(Insurance+PropTax)/12")
    mrow(b + 7, "Site lease", "$k", lambda L, P: f"={L}{b}*{I('N')}/12*{L}$10")
    mrow(b + 8, "Construction capex", "$k", lambda L, P: f"=IF(AND({L}$5>={I('W')},{L}$5<{I('I')}),{I('T')}/{I('H')},0)")
    r = b + 10

section(monthly, r, "Portfolio total", 4)
r += 1
blocks = list(ASSET_BLOCK.values())


def sum_of(offset):
    return lambda L, P: "=" + "+".join(f"{L}{b + offset}" for b in blocks)


mrow(r, "Accredited capacity", "MW", sum_of(2), fmt=K1, total=False, key="acc_mw")
mrow(r + 1, "Capacity revenue", "$k", sum_of(3), key="rev_cap")
mrow(r + 2, "Energy & ancillary margin", "$k", sum_of(4), key="rev_en")
mrow(r + 3, "Total revenue", "$k", lambda L, P: f"={L}{r + 1}+{L}{r + 2}", bold=True, key="rev")
mrow(r + 4, "Fixed O&M", "$k", sum_of(5), key="fom")
mrow(r + 5, "Insurance & property tax", "$k", sum_of(6), key="ins")
mrow(r + 6, "Site lease", "$k", sum_of(7), key="lease")
mrow(r + 7, "Portfolio G&A", "$k", lambda L, P: f"=IF(({'+'.join(f'{L}{b}' for b in blocks)})>0,1,0)*GACost/12*{L}$10", key="ga")
mrow(r + 8, "Total operating costs", "$k", lambda L, P: f"=SUM({L}{r + 4}:{L}{r + 7})", bold=True, key="opex")
mrow(r + 9, "EBITDA", "$k", lambda L, P: f"={L}{r + 3}-{L}{r + 8}", bold=True, key="ebitda")
mrow(r + 10, "Construction capex", "$k", sum_of(8), bold=True, key="capex")
for rr in (r + 3, r + 8, r + 9):
    for c in range(4, LAST_M + 1):
        monthly.cell(row=rr, column=c).border = top_border
monthly.freeze_panes = monthly[f"{col(FIRST)}6"]

# =====================================================================
# Annual
# =====================================================================
title(annual, "Annual cash flows, debt & returns", "Monthly operations rolled up by calendar year; debt sized and repaid annually. $k.")
annual.column_dimensions["B"].width = 40
annual.column_dimensions["C"].width = 13
annual.column_dimensions["D"].width = 13
LAST_Y = FIRST + MODEL_YEARS - 1
LY = col(LAST_Y)
for c in range(FIRST, LAST_Y + 1):
    annual.column_dimensions[col(c)].width = 10
YR = {}
TL = f"{col(FIRST)}{{r}}:{LY}{{r}}"  # timeline range template


def arow(r_, label, formula, fmt=K, total=True, bold=False, key=None, scalar=None, scalar_fmt=None):
    put(annual, f"B{r_}", label, bold=bold)
    for c in range(FIRST, LAST_Y + 1):
        L, P = col(c), col(c - 1)
        f_ = formula(L, P, c == FIRST)
        if f_ is not None:
            put(annual, f"{L}{r_}", f_, fmt=fmt, bold=bold)
    if total:
        put(annual, f"D{r_}", f"=SUM({TL.format(r=r_)})", fmt=fmt, bold=True)
    if key:
        YR[key] = r_


def sumif_monthly(key):
    rr = ROW[key]
    return lambda L, P, first: f"=SUMIF(Monthly!${col(FIRST)}$6:${LM}$6,{L}$5,Monthly!${col(FIRST)}${rr}:${LM}${rr})"


put(annual, "C5", "Scalar", bold=True)
put(annual, "D5", "Total", bold=True)
arow(5, "Year", lambda L, P, first: "=YEAR(ModelStart)" if first else f"={P}5+1", fmt="0", total=False, bold=True)
header_row(annual, 5, FIRST, LAST_Y)

r = 7
section(annual, r, "Operations", 4)
arow(r + 1, "Capacity revenue", sumif_monthly("rev_cap"), key="rev_cap")
arow(r + 2, "Energy & ancillary margin", sumif_monthly("rev_en"), key="rev_en")
arow(r + 3, "Total revenue", lambda L, P, f: f"={L}{r + 1}+{L}{r + 2}", bold=True, key="rev")
arow(r + 4, "Fixed O&M", sumif_monthly("fom"))
arow(r + 5, "Insurance & property tax", sumif_monthly("ins"))
arow(r + 6, "Site lease", sumif_monthly("lease"))
arow(r + 7, "Portfolio G&A", sumif_monthly("ga"))
arow(r + 8, "Total operating costs", lambda L, P, f: f"=SUM({L}{r + 4}:{L}{r + 7})", bold=True, key="opex")
arow(r + 9, "EBITDA (= pre-tax CFADS)", lambda L, P, f: f"={L}{r + 3}-{L}{r + 8}", bold=True, key="ebitda")
arow(r + 10, "EBITDA margin", lambda L, P, f: f"=IF({L}{r + 3}=0,0,{L}{r + 9}/{L}{r + 3})", fmt=PCT, total=False)

r += 12
section(annual, r, "Capex, ITC & tax depreciation", 4)
arow(r + 1, "Construction capex", sumif_monthly("capex"), bold=True, key="capex")
put(annual, f"C{r + 1}", f"=D{r + 1}", fmt=K)
define("TotalCapex", annual, f"C{r + 1}")
arow(r + 2, "ITC cash received (transfer)", lambda L, P, f: f"=SUMPRODUCT((Inputs!$V${A1}:$V${A2}+1={L}$5)*Inputs!$X${A1}:$X${A2})*ITCTransfer", key="itc")
dep_rows = []
rr = r + 3
for a in ASSETS:
    ir = ASSET_ROW[a["name"]]

    def dep(L, P, f, ir=ir):
        kk = f"({L}$5-Inputs!$V${ir}+1)"
        return (f"=IF(OR({kk}<1,{kk}>16),0,IF(Inputs!$P${ir}=5,INDEX(MACRS_5,{kk}),INDEX(MACRS_15,{kk})))"
                f"*Inputs!$Y${ir}")
    arow(rr, f"Tax depreciation — {a['name']}", dep)
    dep_rows.append(rr)
    rr += 1
arow(rr, "Total tax depreciation", lambda L, P, f: f"=SUM({L}{dep_rows[0]}:{L}{dep_rows[-1]})", bold=True, key="dep")

# ---- Debt ----
r = rr + 2
section(annual, r, "Debt sizing (sculpted to target DSCR, drawn at term conversion)", 4)
put(annual, f"B{r + 1}", "Term conversion year (latest live COD)")
put(annual, f"C{r + 1}", f"=MAX(Inputs!$V${A1}:$V${A2})", fmt="0", bold=True)
define("ConvYear", annual, f"C{r + 1}")
put(annual, f"B{r + 2}", "Senior rate (live)")
put(annual, f"C{r + 2}", "=SrRateBase+RateAdj/10000", fmt=PCT2)
define("SrRate", annual, f"C{r + 2}")
put(annual, f"B{r + 3}", "Mezzanine rate (live)")
put(annual, f"C{r + 3}", "=MzRateBase+RateAdj/10000", fmt=PCT2)
define("MzRate", annual, f"C{r + 3}")
r += 5
arow(r, "Years from conversion", lambda L, P, f: f"={L}$5-ConvYear", fmt="0", total=False, key="t")
arow(r + 1, "Debt service period flag", lambda L, P, f: f"=IF(AND({L}{r}>=1,{L}{r}<=Tenor),1,0)", fmt="0", total=False, key="dflag")
arow(r + 2, "CFADS available for debt service", lambda L, P, f: f"={L}{r + 1}*{L}{YR['ebitda']}", key="cfads")

r += 4
put(annual, f"B{r}", "Senior debt", font=f_sect)
arow(r + 1, "Debt service at target DSCR (uncapped)", lambda L, P, f: f"={L}{YR['cfads']}/SrDSCR")
arow(r + 2, "Discount factor at senior rate", lambda L, P, f: f"=IF({L}{YR['dflag']}=1,1/(1+SrRate)^{L}{YR['t']},0)", fmt="0.0000", total=False)
put(annual, f"B{r + 3}", "Debt capacity from sculpting (PV of DS)")
put(annual, f"C{r + 3}", f"=SUMPRODUCT({TL.format(r=r + 1)},{TL.format(r=r + 2)})", fmt=K)
put(annual, f"B{r + 4}", "Gearing cap (senior max gearing × capex)")
put(annual, f"C{r + 4}", "=SrMaxGear*TotalCapex", fmt=K)
put(annual, f"B{r + 5}", "Senior debt amount", bold=True)
put(annual, f"C{r + 5}", f"=MIN(C{r + 3},C{r + 4})", fmt=K, bold=True, border=box)
define("SrAmt", annual, f"C{r + 5}")
put(annual, f"D{r + 5}", f'=IF(C{r + 4}<C{r + 3},"gearing-capped","DSCR-sized")', font=f_sub)
s = r + 6
arow(s, "Senior debt service", lambda L, P, f: f"=IF(C{r + 3}=0,0,{L}{r + 1}*SrAmt/$C${r + 3})", bold=True, key="sr_ds")
arow(s + 1, "Opening balance", lambda L, P, f: "=0" if f else f"={P}{s + 5}", total=False)
arow(s + 2, "Drawdown", lambda L, P, f: f"=IF({L}$5=ConvYear,SrAmt,0)", key="sr_draw")
arow(s + 3, "Interest", lambda L, P, f: f"={L}{s + 1}*SrRate", key="sr_int")
arow(s + 4, "Principal repayment", lambda L, P, f: f"={L}{s}-{L}{s + 3}")
arow(s + 5, "Closing balance", lambda L, P, f: f"={L}{s + 1}+{L}{s + 2}-{L}{s + 4}", total=False, bold=True, key="sr_bal")

r = s + 7
put(annual, f"B{r}", "Mezzanine debt", font=f_sect)
arow(r + 1, "Debt service to combined DSCR, less senior (uncapped)", lambda L, P, f: f"={L}{YR['cfads']}/CombDSCR-{L}{YR['sr_ds']}")
arow(r + 2, "Discount factor at mezzanine rate", lambda L, P, f: f"=IF({L}{YR['dflag']}=1,1/(1+MzRate)^{L}{YR['t']},0)", fmt="0.0000", total=False)
put(annual, f"B{r + 3}", "Debt capacity from sculpting (PV of DS)")
put(annual, f"C{r + 3}", f"=SUMPRODUCT({TL.format(r=r + 1)},{TL.format(r=r + 2)})", fmt=K)
put(annual, f"B{r + 4}", "Room under total gearing cap")
put(annual, f"C{r + 4}", "=TotMaxGear*TotalCapex-SrAmt", fmt=K)
put(annual, f"B{r + 5}", "Mezzanine debt amount", bold=True)
put(annual, f"C{r + 5}", f"=MAX(0,MIN(C{r + 3},C{r + 4}))", fmt=K, bold=True, border=box)
define("MzAmt", annual, f"C{r + 5}")
s = r + 6
arow(s, "Mezzanine debt service", lambda L, P, f: f"=IF($C${r + 3}<=0,0,{L}{r + 1}*MzAmt/$C${r + 3})", bold=True, key="mz_ds")
arow(s + 1, "Opening balance", lambda L, P, f: "=0" if f else f"={P}{s + 5}", total=False)
arow(s + 2, "Drawdown", lambda L, P, f: f"=IF({L}$5=ConvYear,MzAmt,0)", key="mz_draw")
arow(s + 3, "Interest", lambda L, P, f: f"={L}{s + 1}*MzRate", key="mz_int")
arow(s + 4, "Principal repayment", lambda L, P, f: f"={L}{s}-{L}{s + 3}")
arow(s + 5, "Closing balance", lambda L, P, f: f"={L}{s + 1}+{L}{s + 2}-{L}{s + 4}", total=False, bold=True, key="mz_bal")

r = s + 7
arow(r, "Upfront financing fees", lambda L, P, f: f"=IF({L}$5=ConvYear,SrAmt*SrFee+MzAmt*MzFee,0)", key="fees")
arow(r + 1, "Senior DSCR", lambda L, P, f: f'=IF({L}{YR["dflag"]}=1,{L}{YR["cfads"]}/{L}{YR["sr_ds"]},"")', fmt=X2, total=False, key="dscr_s")
arow(r + 2, "Combined DSCR", lambda L, P, f: f'=IF({L}{YR["dflag"]}=1,{L}{YR["cfads"]}/({L}{YR["sr_ds"]}+{L}{YR["mz_ds"]}),"")', fmt=X2, total=False, key="dscr_c")

# ---- Tax ----
r += 4
section(annual, r, "Income tax", 4)
arow(r + 1, "Taxable income (EBITDA − interest − tax depreciation)", lambda L, P, f: f"={L}{YR['ebitda']}-{L}{YR['sr_int']}-{L}{YR['mz_int']}-{L}{YR['dep']}", key="ti")
arow(r + 2, "Income tax (negative = benefit to sponsor)", lambda L, P, f: f"={L}{r + 1}*TaxRate", key="tax")
arow(r + 3, "Unlevered income tax (for project IRR)", lambda L, P, f: f"=({L}{YR['ebitda']}-{L}{YR['dep']})*TaxRate", key="utax")

# ---- Cash flows ----
r += 5
section(annual, r, "Cash flows & returns", 4)
arow(r + 1, "Project (unlevered, after-tax) cash flow",
     lambda L, P, f: f"=-{L}{YR['capex']}+{L}{YR['ebitda']}-{L}{YR['utax']}+{L}{YR['itc']}", bold=True, key="pcf")
arow(r + 2, "Equity (levered, after-tax) cash flow",
     lambda L, P, f: (f"=-{L}{YR['capex']}+{L}{YR['sr_draw']}+{L}{YR['mz_draw']}-{L}{YR['fees']}+{L}{YR['ebitda']}"
                      f"-{L}{YR['sr_ds']}-{L}{YR['mz_ds']}-{L}{YR['tax']}+{L}{YR['itc']}"), bold=True, key="ecf")
arow(r + 3, "Cumulative equity cash flow", lambda L, P, f: f"={L}{r + 2}" if f else f"={P}{r + 3}+{L}{r + 2}", total=False, key="cum")
for rr in (YR["rev"], YR["ebitda"], YR["pcf"], YR["ecf"]):
    for c in range(4, LAST_Y + 1):
        annual.cell(row=rr, column=c).border = top_border
annual.freeze_panes = annual[f"{col(FIRST)}6"]

ytl = lambda key: f"Annual!${col(FIRST)}${YR[key]}:${LY}${YR[key]}"

# =====================================================================
# Summary
# =====================================================================
title(summary, "Battery Storage Portfolio — Project Finance Summary",
      "Front-of-the-meter storage + backup generation. Live outputs for the active scenario. $k unless stated.")
for c_, w in zip("BCDEFGHIJ", [40, 14, 3, 30, 12, 12, 12, 12, 12]):
    summary.column_dimensions[c_].width = w

put(summary, "B5", "Active scenario", bold=True)
put(summary, "C5", "=Scenarios!D5", font=f_link, bold=True)
put(summary, "E5", "Model checks", bold=True)
put(summary, "F5", '=IF(Checks!C4,"ALL OK","CHECK ERRORS")', bold=True)

section(summary, 7, "Returns & valuation", 3)
OUT_ROWS = [
    ("EqIRR", "Equity IRR (levered, after-tax)", f"=IRR({ytl('ecf')})", PCT2),
    ("CostEquityOut", "Target equity IRR", "=CostEquity", PCT2),
    ("ProjIRR", "Project IRR (unlevered, after-tax)", f"=IRR({ytl('pcf')})", PCT2),
    ("WACC", "WACC (discount rate for project NPV)", "=(SrAmt*SrRate*(1-TaxRate)+MzAmt*MzRate*(1-TaxRate)+(TotalCapex-SrAmt-MzAmt)*CostEquity)/TotalCapex", PCT2),
    ("ProjNPV", "Project NPV @ WACC", f"=NPV(WACC,{ytl('pcf')})", K),
    ("EqNPV", "Equity NPV @ cost of equity", f"=NPV(CostEquity,{ytl('ecf')})", K),
    ("Payback", "Equity payback year (cumulative ≥ 0)", f"=IFERROR(INDEX(Annual!${col(FIRST)}$5:${LY}$5,MATCH(TRUE,INDEX({ytl('cum')}>=0,0),0)),\"n/a\")", "0"),
]
r = 8
for name, label, formula, fmt in OUT_ROWS:
    put(summary, f"B{r}", label)
    put(summary, f"C{r}", formula, fmt=fmt, bold=name in ("EqIRR", "ProjNPV"))
    define(name, summary, f"C{r}")
    r += 1
summary.conditional_formatting.add("C8", CellIsRule(operator="greaterThanOrEqual", formula=["CostEquity"], font=Font(color="00753A", bold=True)))
summary.conditional_formatting.add("C8", CellIsRule(operator="lessThan", formula=["CostEquity"], font=Font(color="C00000", bold=True)))

r += 1
section(summary, r, "Capacity pricing", 3)
put(summary, f"B{r + 1}", "Capacity price in live case ($/kW-month, 2027 $)")
put(summary, f"C{r + 1}", "=CapPrice*(1+PriceAdj)", fmt=USD2)
put(summary, f"B{r + 2}", "Required price for target equity IRR", bold=True)
put(summary, f"C{r + 2}", "=RequiredPrice", fmt=USD2, bold=True, font=f_link, border=box)
put(summary, f"B{r + 3}", "Headroom vs. required price")
put(summary, f"C{r + 3}", f"=IFERROR(C{r + 1}/C{r + 2}-1,\"n/a\")", fmt=PCT)
put(summary, f"B{r + 4}", "Required price is interpolated from the Scenarios solver table; Goal Seek EqIRR → CostEquity by changing Inputs!CapPrice gives the same answer.", font=f_sub)

r += 6
section(summary, r, "Debt metrics", 3)
DEBT_ROWS = [
    ("MinDSCRSr", "Minimum senior DSCR", f"=MIN({ytl('dscr_s')})", X2),
    ("AvgDSCRSr", "Average senior DSCR", f"=AVERAGE({ytl('dscr_s')})", X2),
    ("MinDSCRComb", "Minimum combined DSCR", f"=MIN({ytl('dscr_c')})", X2),
    ("Gearing", "Total gearing (debt / capex)", "=(SrAmt+MzAmt)/TotalCapex", PCT),
    ("ConvYearOut", "Term conversion year", "=ConvYear", "0"),
]
for name, label, formula, fmt in DEBT_ROWS:
    r += 1
    put(summary, f"B{r}", label)
    put(summary, f"C{r}", formula, fmt=fmt)
    define(name, summary, f"C{r}")

# Sources & uses + portfolio on the right
for c_ in "EFGHI":
    summary[f"{c_}7"].fill = fill_sect
put(summary, "E7", "Sources & uses at term conversion", font=f_sect)
put(summary, "F8", "$k", bold=True)
put(summary, "G8", "% total", bold=True)
SU = [
    ("Senior debt", "=SrAmt"), ("Mezzanine debt", "=MzAmt"), ("Sponsor equity", "=TotalCapex-SrAmt-MzAmt"),
]
rr = 9
for label, formula in SU:
    put(summary, f"E{rr}", label)
    put(summary, f"F{rr}", formula, fmt=K)
    put(summary, f"G{rr}", f"=F{rr}/$F$12", fmt=PCT)
    rr += 1
put(summary, "E12", "Total sources", bold=True)
put(summary, "F12", "=SUM(F9:F11)", fmt=K, bold=True, border=top_border)
put(summary, "G12", "=SUM(G9:G11)", fmt=PCT, bold=True, border=top_border)
put(summary, "E13", "Total capex (uses)", bold=True)
put(summary, "F13", "=TotalCapex", fmt=K, bold=True)
put(summary, "E14", "ITC cash (transfer, post-COD)")
put(summary, "F14", f"=Annual!D{YR['itc']}", fmt=K)
put(summary, "E15", "Upfront financing fees (equity-funded)")
put(summary, "F15", f"=Annual!D{YR['fees']}", fmt=K)

put(summary, "E17", "Portfolio", font=f_sect)
for c_ in "EFGHI":
    summary[f"{c_}17"].fill = fill_sect
for c_, txt in zip("EFGHI", ["Asset", "MW", "MWh", "Capex $k", "Live COD"]):
    put(summary, f"{c_}18", txt)
header_row(summary, 18, 5, 9)
rr = 19
for a in ASSETS:
    ir = ASSET_ROW[a["name"]]
    put(summary, f"E{rr}", f'=Inputs!B{ir}&" ("&Inputs!C{ir}&")"', font=f_link)
    put(summary, f"F{rr}", f"=Inputs!D{ir}", fmt="#,##0", font=f_link)
    put(summary, f"G{rr}", f"=Inputs!S{ir}", fmt="#,##0", font=f_link)
    put(summary, f"H{rr}", f"=Inputs!T{ir}", fmt=K, font=f_link)
    put(summary, f"I{rr}", f"=Inputs!U{ir}", fmt=DATE, font=f_link)
    rr += 1
put(summary, f"E{rr}", "Total", bold=True)
for c_ in "FGH":
    put(summary, f"{c_}{rr}", f"=SUM({c_}19:{c_}{rr - 1})", fmt=K if c_ == "H" else "#,##0", bold=True, border=top_border)

# Charts
cats = Reference(annual, min_col=FIRST, max_col=LAST_Y, min_row=5)
ch = BarChart()
ch.type = "col"
ch.grouping = "stacked"
ch.overlap = 100
ch.title = "EBITDA vs. debt service ($k)"
ch.height, ch.width = 7.5, 18
for key, label in [("sr_ds", "Senior DS"), ("mz_ds", "Mezz DS")]:
    ref = Reference(annual, min_col=FIRST, max_col=LAST_Y, min_row=YR[key])
    ch.add_data(ref, from_rows=True, titles_from_data=False)
for s_, label in zip(ch.series, ["Senior debt service", "Mezz debt service"]):
    s_.tx = SeriesLabel(v=label)
ch.series[0].graphicalProperties.solidFill = NAVY
ch.series[1].graphicalProperties.solidFill = "C9A227"
ln = LineChart()
ln.add_data(Reference(annual, min_col=FIRST, max_col=LAST_Y, min_row=YR["ebitda"]), from_rows=True, titles_from_data=False)
ln.series[0].tx = SeriesLabel(v="EBITDA (CFADS)")
ln.series[0].graphicalProperties.line.solidFill = TEAL
ln.series[0].graphicalProperties.line.width = 28000
ln.series[0].smooth = False  # openpyxl smooths by default, which invents wiggles between annual points
ch += ln
ch.set_categories(cats)
ch.x_axis.number_format = "0"
ch.x_axis.tickLblSkip = 3
ch.legend.position = "b"
summary.add_chart(ch, "B30")

ch2 = LineChart()
ch2.title = "Debt balances ($k)"
ch2.height, ch2.width = 7.5, 14
for key, label, color in [("sr_bal", "Senior balance", NAVY), ("mz_bal", "Mezz balance", "C9A227")]:
    ch2.add_data(Reference(annual, min_col=FIRST, max_col=LAST_Y, min_row=YR[key]), from_rows=True, titles_from_data=False)
    ch2.series[-1].tx = SeriesLabel(v=label)
    ch2.series[-1].graphicalProperties.line.solidFill = color
    ch2.series[-1].smooth = False
ch2.set_categories(cats)
ch2.legend.position = "b"
ch2.x_axis.tickLblSkip = 3
for chart in (ch, ln, ch2):  # openpyxl 3.1 marks axes as deleted unless told otherwise
    chart.x_axis.delete = False
    chart.y_axis.delete = False
summary.add_chart(ch2, "B46")

# =====================================================================
# Scenarios: data tables
# =====================================================================
OUTPUTS = [("Equity IRR", "=EqIRR", PCT2), ("Project IRR", "=ProjIRR", PCT2), ("Project NPV", "=ProjNPV", K),
           ("Equity NPV", "=EqNPV", K), ("Min senior DSCR", "=MinDSCRSr", X2), ("Total debt", "=SrAmt+MzAmt", K)]


def data_table(top, label, input_cell, values, value_fmt, note=None, outputs=OUTPUTS, left_col=2):
    """One-variable, column-oriented Excel data table. Returns the row after the table."""
    section(scen, top, label, left_col + len(outputs))
    if note:
        put(scen, f"{col(left_col + len(outputs) + 1)}{top}", note, font=f_sub)
    hr = top + 1
    put(scen, f"{col(left_col)}{hr}", f"Input ↓  /  output →")
    for j, (h, _, _) in enumerate(outputs):
        put(scen, f"{col(left_col + 1 + j)}{hr}", h)
    header_row(scen, hr, left_col, left_col + len(outputs))
    fr = hr + 1  # formula row
    put(scen, f"{col(left_col)}{fr}", "Live case →", font=f_sub)
    for j, (_, formula, fmt) in enumerate(outputs):
        put(scen, f"{col(left_col + 1 + j)}{fr}", formula, fmt=fmt, font=f_sub)
    for i, v in enumerate(values):
        inp(scen, f"{col(left_col)}{fr + 1 + i}", v, value_fmt)
        for j, (_, _, fmt) in enumerate(outputs):
            scen[f"{col(left_col + 1 + j)}{fr + 1 + i}"].number_format = fmt
            scen[f"{col(left_col + 1 + j)}{fr + 1 + i}"].font = f_norm
    ref = f"{col(left_col + 1)}{fr + 1}:{col(left_col + len(outputs))}{fr + len(values)}"
    scen[ref.split(":")[0]] = DataTableFormula(ref=ref, dt2D=False, dtr=False, r1=input_cell)
    return fr + len(values) + 2


r = 16
r = data_table(r, "Scenario comparison (varies C5: 1 Base · 2 Downside · 3 Upside)", "C5", [1, 2, 3], "0")
r = data_table(r, "Capex sensitivity (flex on G9)", f"G{DRIVER_ROW['CapexAdj']}", [-0.2, -0.1, -0.05, 0, 0.05, 0.1, 0.2], PCT)
r = data_table(r, "Capacity price sensitivity (flex on G10)", f"G{DRIVER_ROW['PriceAdj']}", [-0.2, -0.1, -0.05, 0, 0.05, 0.1, 0.2], PCT)
r = data_table(r, "Construction delay (flex on G11, months)", f"G{DRIVER_ROW['Delay']}", [0, 3, 6, 9, 12], "0")
r = data_table(r, "Interest rates (flex on G12, bps)", f"G{DRIVER_ROW['RateAdj']}", [-100, -50, 0, 50, 100, 200], "0")

# 2-D: capex flex (rows) x price flex (columns) -> equity IRR
section(scen, r, "Equity IRR: capex flex (rows) × capacity price flex (columns)", 8)
corner = r + 1
put(scen, f"B{corner}", "=EqIRR", fmt=PCT2, font=f_sub)
price_vals = [-0.2, -0.1, 0, 0.1, 0.2]
capex_vals = [-0.2, -0.1, 0, 0.1, 0.2]
for j, v in enumerate(price_vals):
    inp(scen, f"{col(3 + j)}{corner}", v, PCT)
for i, v in enumerate(capex_vals):
    inp(scen, f"B{corner + 1 + i}", v, PCT)
    for j in range(len(price_vals)):
        scen[f"{col(3 + j)}{corner + 1 + i}"].number_format = PCT2
ref2 = f"C{corner + 1}:{col(2 + len(price_vals))}{corner + len(capex_vals)}"
scen[f"C{corner + 1}"] = DataTableFormula(ref=ref2, dt2D=True, r1=f"G{DRIVER_ROW['PriceAdj']}", r2=f"G{DRIVER_ROW['CapexAdj']}")
scen.conditional_formatting.add(ref2, CellIsRule(operator="greaterThanOrEqual", formula=["CostEquity"], fill=PatternFill("solid", fgColor="D9EAD3")))
scen.conditional_formatting.add(ref2, CellIsRule(operator="lessThan", formula=["CostEquity"], fill=PatternFill("solid", fgColor="F4CCCC")))
put(scen, f"H{corner}", "Green = meets target equity IRR", font=f_sub)
r = corner + len(capex_vals) + 2

# Required capacity price solver (right-hand side, columns J-L)
put(scen, "J5", "Required capacity price solver", font=f_sect)
put(scen, "J6", "Equity IRR across capacity-price flex (Data Table on G10); price where IRR crosses the target is interpolated.", font=f_sub)
put(scen, "J7", "Price flex")
put(scen, "K7", "Equity IRR")
put(scen, "L7", "Price $/kW-mo")
header_row(scen, 7, 10, 12)
put(scen, "J8", "Live case →", font=f_sub)
put(scen, "K8", "=IFERROR(EqIRR,-1)", fmt=PCT2, font=f_sub)
flexes = [round(-0.4 + 0.02 * i, 2) for i in range(41)]
s1 = 9
for i, v in enumerate(flexes):
    inp(scen, f"J{s1 + i}", v, PCT)
    scen[f"K{s1 + i}"].number_format = PCT2
    put(scen, f"L{s1 + i}", f"=CapPrice*(1+$F${DRIVER_ROW['PriceAdj']}+J{s1 + i})", fmt=USD2)
s2 = s1 + len(flexes) - 1
scen[f"K{s1}"] = DataTableFormula(ref=f"K{s1}:K{s2}", dt2D=False, dtr=False, r1=f"G{DRIVER_ROW['PriceAdj']}")
J, Kc = f"$J${s1}:$J${s2}", f"$K${s1}:$K${s2}"
put(scen, "N7", "Solver output", bold=True)
put(scen, "N8", "First row ≥ target")
put(scen, "O8", f"=IFERROR(MATCH(TRUE,INDEX({Kc}>=CostEquity,0),0),0)", fmt="0")
put(scen, "N9", "Required price flex")
put(scen, "O9", (f'=IF(OR(O8<=1),NA(),INDEX({J},O8-1)+(CostEquity-INDEX({Kc},O8-1))'
                 f'*(INDEX({J},O8)-INDEX({J},O8-1))/(INDEX({Kc},O8)-INDEX({Kc},O8-1)))'), fmt=PCT2)
put(scen, "N10", "Required capacity price", bold=True)
put(scen, "O10", f'=IFERROR(CapPrice*(1+$F${DRIVER_ROW["PriceAdj"]}+O9),"outside ±40%")', fmt=USD2, bold=True, border=box)
define("RequiredPrice", scen, "O10")
scen.freeze_panes = "A5"

# =====================================================================
# Assumption log (values link live to Inputs)
# =====================================================================
title(log, "Assumption Log", "Every input, who owns it, the basis used, and what to validate it against. Values link live to Inputs.")
LOG_COLS = [("ID", 6), ("Category", 16), ("Assumption", 36), ("Value", 11), ("Unit", 13), ("Owner (source team)", 20),
            ("Basis / rationale", 58), ("Validate against", 50), ("Status", 18)]
for i, (h, w) in enumerate(LOG_COLS):
    put(log, f"{col(2 + i)}5", h)
    log.column_dimensions[col(2 + i)].width = w
header_row(log, 5, 2, 1 + len(LOG_COLS))
fmt_of = {n: f for n, _, _, f, _, _ in GLOBAL_ROWS + FIN_ROWS}
unit_of = {n: u for n, _, _, _, u, _ in GLOBAL_ROWS + FIN_ROWS}
label_of = {n: l for n, l, _, _, _, _ in GLOBAL_ROWS + FIN_ROWS}
LOG = [
    ("Revenue", "CapPrice", "Commercial", "Illustrative 2027 capacity price, set near the level that clears the target return",
     "Recent ISO/RTO capacity auction results (e.g. PJM BRA, NYISO ICAP, ISO-NE FCA) and bilateral RA contract quotes"),
    ("Revenue", "PriceEsc", "Commercial", "Roughly inflation-linked", "Contract escalators in RA/capacity agreements; forward curves"),
    ("Opex", "OMEsc", "Asset Management", "Slightly above CPI for labor-heavy O&M", "O&M / LTSA contract escalators"),
    ("Capex", "Contingency", "Engineering / EPC", "Owner's costs, interconnection true-ups and contingency", "EPC bid spread; interconnection study cost estimates"),
    ("Opex", "Insurance", "Risk / Insurance", "Typical % of replacement cost for BESS", "Broker indications"),
    ("Opex", "PropTax", "Tax", "Varies by jurisdiction and abatements", "County assessor / PILOT agreements"),
    ("Opex", "GACost", "Finance", "Portfolio-level asset management and corporate overhead", "Internal budget"),
    ("Tax", "TaxRate", "Tax", "21% federal + ~5% blended state", "Tax department; state apportionment"),
    ("Tax", "ITCRate", "Tax", "30% base credit assuming prevailing wage & apprenticeship; no adders", "Current §48E rules incl. any phase-down and FEOC restrictions"),
    ("Tax", "ITCTransfer", "Tax / Treasury", "Credit sold to a third party under transferability", "Recent transfer market quotes from advisors / brokers"),
    ("Tax", "ITCBasisRed", "Tax", "Basis reduced by 50% of ITC", "Tax counsel"),
    ("Financing", "SrRateBase", "Treasury / Capital Markets", "Base rate + margin for contracted storage debt", "Lender term sheets; SOFR swap curve"),
    ("Financing", "SrDSCR", "Treasury / Capital Markets", "Contracted capacity revenue supports a mid-1.3x DSCR", "Lender term sheets / rating agency criteria"),
    ("Financing", "SrMaxGear", "Treasury / Capital Markets", "Typical leverage ceiling", "Lender term sheets"),
    ("Financing", "SrFee", "Treasury / Capital Markets", "Arrangement fee", "Lender term sheets"),
    ("Financing", "MzRateBase", "Treasury / Capital Markets", "Subordinated / holdco debt pricing", "Mezz lender indications"),
    ("Financing", "CombDSCR", "Treasury / Capital Markets", "Mezz sized to the gap between senior and combined DSCR", "Mezz lender indications"),
    ("Financing", "TotMaxGear", "Treasury / Capital Markets", "Total leverage ceiling", "Lender term sheets"),
    ("Financing", "MzFee", "Treasury / Capital Markets", "Arrangement fee", "Mezz lender indications"),
    ("Financing", "Tenor", "Treasury / Capital Markets", "Shorter than 20-year asset life", "Lender term sheets; contract tenor"),
    ("Returns", "CostEquity", "Investment Committee", "Sponsor hurdle for contracted storage", "IC hurdle rates; comparable transactions"),
]
r = 6
for i, (cat, name, owner, basis, validate) in enumerate(LOG, start=1):
    vals = [f"A-{i:02d}", cat, label_of[name], f"={name}", unit_of[name], owner, basis, validate, "Illustrative — to validate"]
    for j, v in enumerate(vals):
        c = put(log, f"{col(2 + j)}{r}", v, font=f_link if j == 3 else f_norm, fmt=fmt_of[name] if j == 3 else None)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    r += 1
ASSET_LOG = [
    ("Capex", "F", "EPC unit cost", "Engineering / EPC", "All-in EPC for 4h/2h storage and recip engine backup generation", "EPC bids; NREL Annual Technology Baseline; Lazard LCOE+/LCOS"),
    ("Timing", "I", "COD (base)", "Development / Construction", "Staggered CODs Dec-27 to Jul-28", "Interconnection agreements; EPC schedule"),
    ("Revenue", "J", "Capacity accreditation", "Commercial", "Effective load carrying capability by duration", "ISO accreditation (ELCC) tables"),
    ("Operations", "K", "Degradation / yr", "Engineering / OEM", "No augmentation assumed, so accredited MW falls each year", "OEM warranty curves"),
    ("Opex", "L", "Fixed O&M $/kW-yr", "Asset Management", "Includes LTSA and warranty extension", "O&M / LTSA quotes"),
    ("Revenue", "M", "Energy & ancillary margin $/kW-yr", "Commercial / Trading", "Conservative merchant upside on top of capacity", "Historical arbitrage & AS spreads by node"),
]
for cat, c_, label, owner, basis, validate in ASSET_LOG:
    for a in ASSETS:
        ir = ASSET_ROW[a["name"]]
        fmt = inputs[f"{c_}{ir}"].number_format
        vals = [f"A-{r - 5:02d}", cat, f"{a['name']} — {label}", f"=Inputs!{c_}{ir}",
                {"F": a["unit"], "I": "date", "J": "%", "K": "% / yr", "L": "$/kW-yr", "M": "$/kW-yr"}[c_],
                owner, basis, validate, "Illustrative — to validate"]
        for j, v in enumerate(vals):
            cell = put(log, f"{col(2 + j)}{r}", v, font=f_link if j == 3 else f_norm, fmt=fmt if j == 3 else None)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1
log.freeze_panes = "B6"
log.auto_filter.ref = f"B5:{col(1 + len(LOG_COLS))}{r - 1}"

# =====================================================================
# Checks
# =====================================================================
title(checks, "Model Checks", "Integrity checks; all must be TRUE. The overall flag feeds the Summary sheet.")
checks.column_dimensions["B"].width = 60
checks.column_dimensions["C"].width = 12
checks.column_dimensions["D"].width = 16
put(checks, "B4", "ALL CHECKS PASS", bold=True)
tol = 0.01
CHECKS = [
    ("Senior debt fully repaid by end of tenor", f"ABS(Annual!{LY}{YR['sr_bal']})<{tol}", f"=Annual!{LY}{YR['sr_bal']}"),
    ("Mezzanine debt fully repaid by end of tenor", f"ABS(Annual!{LY}{YR['mz_bal']})<{tol}", f"=Annual!{LY}{YR['mz_bal']}"),
    ("Min senior DSCR ≥ target", "MinDSCRSr>=SrDSCR-0.0001", "=MinDSCRSr"),
    ("Min combined DSCR ≥ target", "MinDSCRComb>=CombDSCR-0.0001", "=MinDSCRComb"),
    ("Sources = uses", "ABS(Summary!F12-Summary!F13)<0.01", "=Summary!F12-Summary!F13"),
    ("Monthly capex = asset capex on Inputs", f"ABS(Monthly!D{ROW['capex']}-TotalCapexInputs)<{tol}", f"=Monthly!D{ROW['capex']}-TotalCapexInputs"),
    ("Annual revenue = monthly revenue", f"ABS(Annual!D{YR['rev']}-Monthly!D{ROW['rev']})<{tol}", f"=Annual!D{YR['rev']}-Monthly!D{ROW['rev']}"),
    ("Annual EBITDA = monthly EBITDA", f"ABS(Annual!D{YR['ebitda']}-Monthly!D{ROW['ebitda']})<{tol}", f"=Annual!D{YR['ebitda']}-Monthly!D{ROW['ebitda']}"),
    ("Total gearing ≤ cap", "Gearing<=TotMaxGear+0.0001", "=Gearing"),
    ("Construction starts on/after model start", f"MIN(Inputs!W{A1}:W{A2})>=ModelStart", f"=MIN(Inputs!W{A1}:W{A2})"),
    ("All asset lives end within model horizon", f"MAX(Inputs!Z{A1}:Z{A2})<=EDATE(ModelStart,{MODEL_YEARS * 12})", f"=MAX(Inputs!Z{A1}:Z{A2})"),
    ("Tax depreciation = depreciable basis", f"ABS(Annual!D{YR['dep']}-Inputs!Y{A2 + 1})<{tol}", f"=Annual!D{YR['dep']}-Inputs!Y{A2 + 1}"),
    ("Scenario selector is 1, 2 or 3", "OR(ScenarioSel=1,ScenarioSel=2,ScenarioSel=3)", "=ScenarioSel"),
]
for c_, txt in zip("BCD", ["Check", "Pass?", "Value"]):
    put(checks, f"{c_}6", txt)
header_row(checks, 6, 2, 4)
for i, (label, cond, val) in enumerate(CHECKS):
    rr = 7 + i
    put(checks, f"B{rr}", label)
    put(checks, f"C{rr}", f"={cond}", bold=True)
    put(checks, f"D{rr}", val, fmt="#,##0.00")
    if "date" in label.lower() or "start" in label.lower() or "horizon" in label.lower():
        checks[f"D{rr}"].number_format = "dd-mmm-yyyy"
last = 7 + len(CHECKS) - 1
put(checks, "C4", f"=AND(C7:C{last})", bold=True, border=box)
for ref in ("C4", f"C7:C{last}"):
    checks.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=["TRUE"], fill=PatternFill("solid", fgColor="D9EAD3")))
    checks.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=["FALSE"], fill=PatternFill("solid", fgColor="F4CCCC")))

# =====================================================================
# Cover
# =====================================================================
title(cover, "Battery Storage Project Finance & Capacity Pricing Model",
      "Portfolio of front-of-the-meter battery storage and backup generation assets · Kshama Upadhyay")
cover.column_dimensions["B"].width = 26
cover.column_dimensions["C"].width = 100
rows = [
    ("Purpose", "Value a 400 MW portfolio (four BESS sites + one backup generation site), size a senior/mezzanine/equity capital stack, "
                "and solve for the capacity price that delivers the sponsor's target equity IRR."),
    ("Status", "All assumptions are illustrative and logged in 'Assumption Log' with an owner and a validation source. Not investment advice."),
    ("Units", "Nominal USD, $ thousands ($k) unless stated. Monthly operating timeline Jan-2027 to Dec-2049; financing and returns annual."),
    ("", ""),
    ("Sheet guide", ""),
    ("Summary", "Returns, capital stack, required capacity price, debt metrics, charts."),
    ("Inputs", "General, market, tax and financing assumptions, asset-level table, MACRS tables."),
    ("Scenarios", "Active-case selector (Base / Downside / Upside), sensitivity flex cells, Excel Data Tables, required-price solver."),
    ("Monthly", "Asset-level operating flags, degradation, revenue, opex and construction capex by month."),
    ("Annual", "Annual roll-up, ITC, MACRS depreciation, sculpted senior and mezz debt, tax, project and equity cash flows."),
    ("Assumption Log", "Every input with owner, basis and validation source; values link live to Inputs."),
    ("Checks", "Integrity checks (debt repaid, DSCR floors, sources = uses, monthly-to-annual tie-outs)."),
    ("", ""),
    ("Conventions", "Blue on yellow = hard-coded input · black = formula · green = link to another sheet. No hard-codes in calculation sheets."),
    ("Methodology", "Debt is sized at term conversion (latest COD) by sculpting debt service to CFADS / target DSCR, capped by gearing. "
                    "Mezzanine takes the gap between the senior and combined DSCR targets. CFADS for sizing is pre-tax EBITDA (no circularity). "
                    "Storage ITC is monetized by transfer the year after COD. Tax losses are assumed usable by the sponsor."),
    ("How to re-solve price", "Scenarios!O10 updates automatically on recalc. Or use Data → What-If Analysis → Goal Seek: set Summary!C8 (Equity IRR) "
                              "to the target by changing Inputs!CapPrice."),
    ("Recalculation", "If Data Tables look stale, press F9 (or set Formulas → Calculation Options → Automatic)."),
]
for i, (k_, v) in enumerate(rows):
    rr = 5 + i
    put(cover, f"B{rr}", k_, bold=True, font=f_sect if k_ in ("Sheet guide",) else f_bold)
    c = put(cover, f"C{rr}", v)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    cover[f"B{rr}"].alignment = Alignment(vertical="top")

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90
# Summary prints on one landscape page
summary.page_setup.orientation = "landscape"
summary.sheet_properties.pageSetUpPr.fitToPage = True
summary.page_setup.fitToWidth = 1
summary.page_setup.fitToHeight = 1
summary.print_area = "A1:J62"
wb.calculation.fullCalcOnLoad = True
wb.calculation.calcMode = "auto"
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"Saved {OUT}")
